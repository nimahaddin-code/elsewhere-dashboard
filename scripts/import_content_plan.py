from __future__ import annotations

import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "outputs" / "elsewhere-60-slot" / "Elsewhere_Co_30_Day_Threads_X.xlsx"
OUTPUT = ROOT / "lib" / "content-plan-data.ts"


def rows_by_header(ws, header_row: int = 6):
    headers = [cell.value for cell in ws[header_row]]
    rows = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not values[0]:
            continue
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        rows.append(row)
    return rows


def serialise(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return value


wb = load_workbook(SOURCE, data_only=True)
master = rows_by_header(wb["30-Day Master Plan"])
threads = {row["Slot ID"]: row for row in rows_by_header(wb["Threads Ready-to-Generate"])}
x_rows = {row["Slot ID"]: row for row in rows_by_header(wb["X Ready-to-Generate"])}
angles = {row["Angle"]: row for row in rows_by_header(wb["Copy Angle Library"])}

slots = []
for row in master:
    slot_id = row["Slot ID"]
    thread = threads[slot_id]
    x_row = x_rows[slot_id]
    angle = angles[row["Angle"]]
    slots.append(
        {
            "id": slot_id,
            "date": serialise(row["Tanggal"]),
            "period": row["Waktu"],
            "time": serialise(row["Jam WITA"]),
            "stage": row["Stage"],
            "angle": row["Angle"],
            "category": row["Kategori"],
            "productId": row["Product ID"],
            "initialStatus": row["Status"],
            "nextAction": row["Aksi berikutnya"],
            "threadsTemplate": thread["Full Copy Template"],
            "threadsAttachment": thread["RAW Image Attachment"],
            "threadsGuidance": thread["Arahan komentar produk"],
            "threadsRequirement": thread["Syarat produk / bukti"],
            "xTemplate": x_row["Full Copy Template"],
            "xAttachment": x_row["RAW Image Attachment"],
            "xGuidance": x_row["Arahan komentar produk"],
            "xOverflowPolicy": x_row["Overflow policy"],
            "angleId": angle["Angle ID"],
            "bridgeThreads": angle["Bridge Threads"],
            "bridgeX": angle["Bridge X"],
            "factGuidance": angle["Arahan fakta produk"],
            "eligibility": angle["Eligibility"],
            "fallback": angle["Fallback"],
        }
    )

assert len(slots) == 60
assert len({slot["id"] for slot in slots}) == 60
assert all(sum(1 for item in slots if item["date"] == slot["date"]) == 2 for slot in slots)

config = {
    "campaignId": "elsewhere-kl-2026",
    "planVersion": "1.0",
    "planStart": "2026-09-17",
    "planEnd": "2026-10-16",
    "timezone": "Asia/Makassar",
    "defaultTimes": {"Siang": "12:00", "Malam": "20:00"},
    "waUrl": "https://chat.whatsapp.com/JxEr1IQIcel8dn0bEu2HQK?mode=gi_t",
    "xMaxLength": 280,
    "threadsMaxLength": 500,
}

payload = json.dumps(slots, ensure_ascii=False, indent=2, default=serialise)
config_payload = json.dumps(config, ensure_ascii=False, indent=2)
OUTPUT.write_text(
    "// Generated from Elsewhere_Co_30_Day_Threads_X.xlsx. Do not hand-edit templates.\n"
    f"export const contentPlanConfig = {config_payload} as const;\n\n"
    f"export const contentPlanSlots = {payload} as const;\n",
    encoding="utf-8",
)
print(f"Imported {len(slots)} slots to {OUTPUT}")
